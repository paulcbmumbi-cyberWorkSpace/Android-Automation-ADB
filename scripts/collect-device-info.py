import shutil
import subprocess
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


OUTPUT_WORKBOOK = Path(__file__).resolve().parent / "device_info.xlsx"


def get_adb_executable() -> str | None:
    """Return the ADB executable found through the system PATH."""
    return shutil.which("adb")


def get_device_name(adb: str, serial: str) -> str:
    """Return the Android model name for a device."""
    try:
        result = subprocess.run(
            [adb, "-s", serial, "shell", "getprop", "ro.product.model"],
            check=True,
            capture_output=True,
            text=True,
        )
        return result.stdout.strip() or "Unknown"
    except subprocess.CalledProcessError:
        return "Unknown"


def collect_device_info(output_path: Path = OUTPUT_WORKBOOK) -> None:
    """Store each authorized device name and serial number in an Excel workbook."""
    adb = get_adb_executable()
    if adb is None:
        print("ADB is not installed or not available in PATH.")
        return

    try:
        result = subprocess.run(
            [adb, "devices"],
            check=True,
            capture_output=True,
            text=True,
        )
    except subprocess.CalledProcessError as error:
        print(f"Failed to list Android devices: {error}")
        return

    devices = []
    for line in result.stdout.splitlines()[1:]:
        columns = line.split()
        if len(columns) >= 2 and columns[1] == "device":
            serial = columns[0]
            devices.append((get_device_name(adb, serial), serial))

    if not devices:
        print("No authorized Android devices found.")
        return

    if output_path.exists():
        workbook = load_workbook(output_path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Device Info"
        worksheet.append(["Satellite Name", "Device Name", "Serial Number", "Collected At"])

    headers = [cell.value for cell in worksheet[1]]
    if headers[:3] == ["Device Name", "Serial Number", "Collected At"]:
        worksheet.insert_cols(1)
        worksheet.cell(row=1, column=1, value="Satellite Name")

    existing_serials = {cell.value for cell in worksheet["C"][1:]}
    for device_name, serial in devices:
        if serial not in existing_serials:
            satellite_name = input(
                f"Enter satellite name for {device_name} ({serial}): "
            ).strip()
            if not satellite_name:
                satellite_name = "Unnamed"
            worksheet.append(
                [
                    satellite_name,
                    device_name,
                    serial,
                    datetime.now().isoformat(timespec="seconds"),
                ]
            )
            existing_serials.add(serial)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(f"Device information stored in {output_path}")


if __name__ == "__main__":
    collect_device_info()
